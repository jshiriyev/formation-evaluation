import contextlib
import datetime

from difflib import SequenceMatcher

import logging

import os
import re

import numpy
import openpyxl

from ._browser import Browser

class XlSheet(Browser):

    def __init__(self,*args,**kwargs):

        super().__init__(*args,**kwargs)

    def write(self,filepath,title):

        wb = openpyxl.Workbook()

        sheet = wb.active

        if title is not None:
            sheet.title = title

        for line in running:
            sheet.append(line)

        wb.save(filepath)

def loadsheet(filepath,sheetname,**kwargs):

    pass

class XlWorm():

    def __init__(self,filepath,homedir=None,filedir=None,**kwargs):

        with xlbatch.xlopen(filepath) as book:

            for sheet in sheets:

                print("Loading {} {}".format(filepath,sheet))

                dataframe = self.load(book,sheet,**kwargs)

                self.frames.append(dataframe)

    def load(self,book,sheet,sheetsearch=False,min_row=1,min_col=1,max_row=None,max_col=None,hrows=0):
        """It reads provided excel worksheet and returns it as a frame."""

        dataframe = frame()

        if sheet is None:
            sheetname = book.sheetnames[0]
        elif isinstance(sheet,int):
            sheetname = book.sheetnames[sheet]
        elif isinstance(sheet,str) and sheetsearch:
            srcscores = [SequenceMatcher(None,page,sheet).ratio() for page in book.sheetnames]
            sheetname = book.sheetnames[srcscores.index(max(srcscores))]
        elif isinstance(sheet,str):
            if sheet in book.sheetnames:
                sheetname = sheet
            else:
                raise ValueError(f"'{sheet}' could not be found in the xlbook, try sheetsearch=True.")
        else:
            raise TypeError(f"Expected sheet is either none, int or str, but the input type is {type(sheet[1])}.")

        # frame.sheetname = sheetname

        rows = book[sheetname].iter_rows(
            min_row=min_row,max_row=max_row,
            min_col=min_col,max_col=max_col,
            values_only=True)

        cols = []
    
        for index,row in enumerate(rows):
            if index==0:
                [cols.append([]) for _ in row]
            if any(row):
                [col.append(cell) for cell,col in zip(row,cols)]

        heads_ = key2column(size=(len(cols)+min_col-1),dtype='str')
        heads_ = heads_.vals[min_col-1:]

        for index,col in enumerate(cols):
            if any(col):
                info = " ".join([str(cell) for cell in col[:hrows] if cell is not None])
                datacolumn = column(col[hrows:],head=heads_[index],info=info)
                dataframe._setup(datacolumn)

        return dataframe

    @contextlib.contextmanager
    def xlopen(filepath):
        xlbook = openpyxl.load_workbook(filepath,read_only=True,data_only=True)
        try:
            yield xlbook
        finally:
            xlbook._archive.close()